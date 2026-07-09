"""
AMAZON AI - Semantic Search (Type B: Content-Based Discovery)
Finds files by understanding their content using embeddings + FAISS vector search.

Model: all-MiniLM-L6-v2 (22M params, 384-dim vectors)
Index: ml/semantic_index/ (vectors.faiss + metadata.json)
"""
import os
import json
import logging
from pathlib import Path
from datetime import datetime
from typing import List, Dict, Optional

import numpy as np

# Index storage location
_INDEX_DIR = Path(__file__).parent.parent / "ml" / "semantic_index"
_INDEX_FILE = _INDEX_DIR / "vectors.faiss"
_META_FILE = _INDEX_DIR / "metadata.json"
_MODEL_NAME = "all-MiniLM-L6-v2"

# Directories to scan for indexing
_SCAN_DIRS = [
    Path.home() / "Desktop",
    Path.home() / "Documents",
    Path.home() / "Downloads",
]

# Supported file extensions for content extraction
_SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".doc", ".pptx", ".ppt",
    ".txt", ".md", ".csv", ".json", ".xml", ".rtf",
    ".xlsx", ".xls", ".py", ".js", ".html", ".css",
}

# Skip patterns
_SKIP_DIRS = {
    "__pycache__", ".git", "node_modules", ".vscode", ".idea",
    "venv", "env", ".venv", "temp", "tmp", "AppData",
}

# Max characters to extract for embedding (avoids memory issues)
_MAX_CHARS = 2000

# Lazy-loaded globals
_model = None
_faiss = None

logger = logging.getLogger(__name__)


# =====================================================================
#  LAZY LOADING (model + FAISS loaded only when needed)
# =====================================================================

def _get_model():
    """Lazy-load the sentence-transformer model (cached after first call)."""
    global _model
    if _model is None:
        print(f"Loading semantic model: {_MODEL_NAME}...")
        from sentence_transformers import SentenceTransformer
        _model = SentenceTransformer(_MODEL_NAME)
    return _model


def _get_faiss():
    """Lazy-load FAISS library."""
    global _faiss
    if _faiss is None:
        import faiss as faiss_lib
        _faiss = faiss_lib
    return _faiss


# =====================================================================
#  CONTENT EXTRACTION
# =====================================================================

def extract_text(file_path: Path) -> Optional[str]:
    """
    Extract text content from a file for embedding.
    Returns first ~2000 chars of text, or None if extraction fails.
    """
    ext = file_path.suffix.lower()
    try:
        if ext == ".pdf":
            return _extract_pdf(file_path)
        elif ext in (".docx", ".doc"):
            return _extract_docx(file_path)
        elif ext in (".pptx", ".ppt"):
            return _extract_pptx(file_path)
        elif ext in (".xlsx", ".xls"):
            return _extract_xlsx(file_path)
        elif ext in (".txt", ".md", ".csv", ".json", ".xml", ".rtf", ".py", ".js", ".html", ".css"):
            return _extract_text_file(file_path)
    except Exception as e:
        logger.debug(f"Failed to extract text from {file_path}: {e}")
    return None


def _extract_pdf(path: Path) -> Optional[str]:
    """Extract text from PDF using PyPDF2."""
    from PyPDF2 import PdfReader
    reader = PdfReader(str(path))
    text_parts = []
    total_chars = 0
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
            total_chars += len(page_text)
            if total_chars >= _MAX_CHARS:
                break
    return "\n".join(text_parts)[:_MAX_CHARS] if text_parts else None


def _extract_docx(path: Path) -> Optional[str]:
    """Extract text from DOCX using python-docx."""
    from docx import Document
    doc = Document(str(path))
    text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(text_parts)[:_MAX_CHARS] if text_parts else None


def _extract_pptx(path: Path) -> Optional[str]:
    """Extract text from PPTX using python-pptx."""
    from pptx import Presentation
    prs = Presentation(str(path))
    text_parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text_parts.append(shape.text)
    return "\n".join(text_parts)[:_MAX_CHARS] if text_parts else None


def _extract_xlsx(path: Path) -> Optional[str]:
    """Extract text from XLSX using openpyxl."""
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    text_parts = []
    total_chars = 0
    for sheet in wb.sheetnames[:3]:  # First 3 sheets only
        ws = wb[sheet]
        for row in ws.iter_rows(max_row=50, values_only=True):  # First 50 rows
            cells = [str(c) for c in row if c is not None]
            if cells:
                line = " ".join(cells)
                text_parts.append(line)
                total_chars += len(line)
                if total_chars >= _MAX_CHARS:
                    break
        if total_chars >= _MAX_CHARS:
            break
    wb.close()
    return "\n".join(text_parts)[:_MAX_CHARS] if text_parts else None


def _extract_text_file(path: Path) -> Optional[str]:
    """Extract text from plain text files."""
    encodings = ["utf-8", "latin-1", "cp1252"]
    for enc in encodings:
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read(_MAX_CHARS)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return None


# =====================================================================
#  FILE SCANNING
# =====================================================================

def _scan_files(search_dirs=None):
    """
    Scan directories for indexable files.

    Returns:
        list of Path objects for files that can be extracted
    """
    dirs = search_dirs or _SCAN_DIRS
    files = []

    for d in dirs:
        if not d.exists():
            continue
        try:
            for item in d.rglob("*"):
                if item.is_dir():
                    continue
                if any(skip in item.parts for skip in _SKIP_DIRS):
                    continue
                if item.suffix.lower() in _SUPPORTED_EXTENSIONS:
                    files.append(item)
        except PermissionError:
            continue

    return files


def _get_file_mtime(path: Path) -> float:
    """Get file modification time."""
    try:
        return path.stat().st_mtime
    except OSError:
        return 0.0


# =====================================================================
#  FAISS INDEX MANAGEMENT
# =====================================================================

def is_index_ready() -> bool:
    """Check if a semantic index exists and is usable."""
    return _INDEX_FILE.exists() and _META_FILE.exists()


def load_index():
    """
    Load existing FAISS index and metadata from disk.

    Returns:
        (faiss_index, metadata_list) or (None, []) if not found
    """
    if not is_index_ready():
        return None, []

    faiss = _get_faiss()
    try:
        index = faiss.read_index(str(_INDEX_FILE))
        with open(_META_FILE, "r", encoding="utf-8") as f:
            metadata = json.load(f)
        return index, metadata
    except Exception as e:
        logger.error(f"Failed to load index: {e}")
        return None, []


def save_index(index, metadata: list):
    """Save FAISS index and metadata to disk."""
    _INDEX_DIR.mkdir(parents=True, exist_ok=True)
    faiss = _get_faiss()
    faiss.write_index(index, str(_INDEX_FILE))
    with open(_META_FILE, "w", encoding="utf-8") as f:
        json.dump(metadata, f, indent=2, ensure_ascii=False)


def build_index(search_dirs=None, force=False) -> dict:
    """
    Build or update the semantic index.

    Args:
        search_dirs: directories to scan (default: Desktop, Documents, Downloads)
        force: if True, rebuild from scratch (ignore existing index)

    Returns:
        dict with stats: {"indexed": int, "skipped": int, "failed": int, "total": int}
    """
    faiss = _get_faiss()
    model = _get_model()

    # Scan files
    files = _scan_files(search_dirs)

    # Load existing index for incremental updates
    if not force and is_index_ready():
        existing_index, existing_meta = load_index()
        meta_by_path = {m["path"]: m for m in existing_meta}
    else:
        existing_index = None
        meta_by_path = {}

    # Extract text from new/changed files
    texts = []
    file_paths = []
    stats = {"indexed": 0, "skipped": 0, "failed": 0, "total": len(files)}

    for fp in files:
        fp_str = str(fp.resolve())
        mtime = _get_file_mtime(fp)

        # Skip if already indexed and not modified
        if not force and fp_str in meta_by_path:
            old_mtime = meta_by_path[fp_str].get("mtime", 0)
            if abs(mtime - old_mtime) < 1.0:  # Within 1 second
                stats["skipped"] += 1
                continue

        # Extract text
        text = extract_text(fp)
        if text and len(text.strip()) > 20:  # Skip near-empty files
            texts.append(text)
            file_paths.append((fp_str, mtime))
            stats["indexed"] += 1
        else:
            stats["failed"] += 1

    # Build new vectors
    if texts:
        print(f"Embedding {len(texts)} documents...")
        embeddings = model.encode(texts, show_progress_bar=False, batch_size=32)
        embeddings = np.array(embeddings, dtype="float32")

        # Normalize for cosine similarity
        faiss.normalize_L2(embeddings)

        # Create or extend index
        dim = embeddings.shape[1]
        if existing_index is None or force:
            index = faiss.IndexFlatIP(dim)  # Inner product (cosine after normalization)
            new_meta = []
        else:
            index = existing_index
            new_meta = list(meta_by_path.values())

        # Add new vectors
        index.add(embeddings)

        # Build metadata for new entries
        for fp_str, mtime in file_paths:
            new_meta.append({
                "path": fp_str,
                "name": Path(fp_str).name,
                "mtime": mtime,
                "indexed_at": datetime.now().isoformat(),
            })

        save_index(index, new_meta)
        stats["total_vectors"] = index.ntotal
    elif not force:
        # No new files, keep existing index
        stats["total_vectors"] = existing_index.ntotal if existing_index else 0
    else:
        stats["total_vectors"] = 0

    print(f"Index complete: {stats['indexed']} indexed, {stats['skipped']} skipped, {stats['failed']} failed")
    return stats


# =====================================================================
#  SEMANTIC SEARCH
# =====================================================================

def search(query: str, top_k: int = 5) -> List[Dict]:
    """
    Semantic search: find files whose content matches the query.

    Args:
        query: natural language query (e.g. "proposal about AI")
        top_k: number of results to return

    Returns:
        list of dicts: [{"name", "path", "score", "modified"}]
    """
    if not is_index_ready():
        return []

    index, metadata = load_index()
    if index is None or index.ntotal == 0:
        return []

    model = _get_model()
    faiss = _get_faiss()

    # Embed query
    query_vec = model.encode([query], show_progress_bar=False)
    query_vec = np.array(query_vec, dtype="float32")
    faiss.normalize_L2(query_vec)

    # Search
    k = min(top_k, index.ntotal)
    scores, indices = index.search(query_vec, k)

    # Map results to file metadata
    results = []
    for score, idx in zip(scores[0], indices[0]):
        if idx < 0 or idx >= len(metadata):
            continue
        meta = metadata[idx]
        mtime = meta.get("mtime", 0)
        modified_str = datetime.fromtimestamp(mtime).strftime("%Y-%m-%d %H:%M") if mtime else "unknown"
        results.append({
            "name": meta.get("name", Path(meta["path"]).name),
            "path": meta["path"],
            "score": round(float(score), 4),
            "modified": modified_str,
        })

    return results
