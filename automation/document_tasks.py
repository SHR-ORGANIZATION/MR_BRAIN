import os
from pathlib import Path
from datetime import datetime

try:
    import docx
    from docx.shared import Pt
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl
    from openpyxl import Workbook
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False


def _build_result(action, path, status, message, item_type="Document", extra=None):
    return {
        "action": action,
        "name": Path(path).name if path else "",
        "path": str(Path(path).expanduser().resolve()) if path else "",
        "type": item_type,
        "status": status,
        "message": message,
        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "details": extra,
    }


def _ensure_parent(path):
    path = Path(path).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def create_text_document(path, content=""):
    path = _ensure_parent(path)
    try:
        with open(path, "w", encoding="utf-8") as file:
            file.write(content or "")
        return _build_result("Created Document", path, "success", f"Text file created: {path}", "Text Document")
    except Exception as e:
        return _build_result("Create Document", path, "failed", str(e), "Text Document")


def create_structured_word_document(doc, path):
    """Save a pre-built docx.Document object to file."""
    path = _ensure_parent(path)
    if not DOCX_AVAILABLE:
        return _build_result(
            "Create Document",
            path,
            "failed",
            "python-docx is not installed. Install it to create Word documents.",
            "Word Document",
        )
    try:
        doc.save(str(path))
        return _build_result("Created Document", path, "success", f"Word document created: {path}", "Word Document")
    except Exception as e:
        return _build_result("Create Document", path, "failed", str(e), "Word Document")


def create_word_document(path, content="", structured=False):
    path = _ensure_parent(path)
    if not DOCX_AVAILABLE:
        return _build_result(
            "Create Document",
            path,
            "failed",
            "python-docx is not installed. Install it to create Word documents.",
            "Word Document",
        )
    try:
        doc = docx.Document()
        if structured and isinstance(content, dict):
            for section_title, section_content in content.items():
                doc.add_heading(section_title, level=1)
                for paragraph in str(section_content).split("\n"):
                    p = doc.add_paragraph(paragraph)
                    if p.runs:
                        p_format = p.runs[0].font
                        p_format.size = Pt(12)
        else:
            for paragraph in str(content).split("\n"):
                p = doc.add_paragraph(paragraph)
                if p.runs:
                    p_format = p.runs[0].font
                    p_format.size = Pt(12)
        doc.save(path)
        return _build_result("Created Document", path, "success", f"Word document created: {path}", "Word Document")
    except Exception as e:
        return _build_result("Create Document", path, "failed", str(e), "Word Document")


def create_excel_document(path, content=None):
    path = _ensure_parent(path)
    if not XLSX_AVAILABLE:
        return _build_result(
            "Create Document",
            path,
            "failed",
            "openpyxl is not installed. Install it to create Excel spreadsheets.",
            "Excel Spreadsheet",
        )
    try:
        workbook = Workbook()
        sheet = workbook.active
        if isinstance(content, list):
            for row_idx, row in enumerate(content, start=1):
                if isinstance(row, (list, tuple)):
                    for col_idx, value in enumerate(row, start=1):
                        sheet.cell(row=row_idx, column=col_idx, value=str(value))
                else:
                    sheet.cell(row=row_idx, column=1, value=str(row))
        else:
            sheet.cell(row=1, column=1, value=str(content or ""))
        workbook.save(path)
        return _build_result("Created Document", path, "success", f"Excel spreadsheet created: {path}", "Excel Spreadsheet")
    except Exception as e:
        return _build_result("Create Document", path, "failed", str(e), "Excel Spreadsheet")


def create_ppt_document(path, content=""):
    path = _ensure_parent(path)
    if not PPTX_AVAILABLE:
        return _build_result(
            "Create Document",
            path,
            "failed",
            "python-pptx is not installed. Install it to create PowerPoint presentations.",
            "PowerPoint Presentation",
        )
    try:
        presentation = Presentation()
        title_slide_layout = presentation.slide_layouts[1] if len(presentation.slide_layouts) > 1 else presentation.slide_layouts[0]
        slide = presentation.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        body = slide.shapes.placeholders[1]
        title.text = Path(path).stem
        tf = body.text_frame
        for paragraph in str(content).split("\n"):
            p = tf.add_paragraph()
            p.text = paragraph
            p.level = 0
        presentation.save(path)
        return _build_result("Created Document", path, "success", f"PowerPoint presentation created: {path}", "PowerPoint Presentation")
    except Exception as e:
        return _build_result("Create Document", path, "failed", str(e), "PowerPoint Presentation")


def create_pdf_document(path, content=""):
    path = _ensure_parent(path)
    if not REPORTLAB_AVAILABLE:
        return _build_result(
            "Create Document",
            path,
            "failed",
            "reportlab is not installed. Install it to create PDF documents.",
            "PDF Document",
        )
    try:
        c = canvas.Canvas(str(path), pagesize=letter)
        width, height = letter
        lines = str(content).split("\n")
        y = height - 72
        for line in lines:
            c.drawString(72, y, line)
            y -= 18
            if y < 72:
                c.showPage()
                y = height - 72
        c.save()
        return _build_result("Created Document", path, "success", f"PDF document created: {path}", "PDF Document")
    except Exception as e:
        return _build_result("Create Document", path, "failed", str(e), "PDF Document")


def create_document(path, content="", topic=None):
    ext = Path(path).suffix.lower()
    if not ext:
        return _build_result("Create Document", path, "failed", "No document extension provided", "Document")
    if ext in [".txt", ".md", ".csv"]:
        return create_text_document(path, content)
    if ext == ".docx":
        return create_word_document(path, content)
    if ext in [".xlsx", ".xlsm", ".xltx"]:
        rows = [str(content).split("\n")] if content else []
        return create_excel_document(path, rows)
    if ext in [".pptx"]:
        return create_ppt_document(path, content)
    if ext == ".pdf":
        return create_pdf_document(path, content)
    return _build_result("Create Document", path, "failed", f"Unsupported document type: {ext}", "Document")


def generate_document(path, topic, content=None, doc_type=None, instruction=None):
    import importlib
    import automation.document_generator as document_generator

    # Reload in-process so UI sessions pick up template changes without restart
    document_generator = importlib.reload(document_generator)
    generate_assignment = document_generator.generate_assignment
    generate_research_proposal = document_generator.generate_research_proposal
    generate_report = document_generator.generate_report
    generate_schedule = document_generator.generate_schedule
    generate_dynamic_document = document_generator.generate_dynamic_document
    generate_exam_document = document_generator.generate_exam_document
    save_document = document_generator.save_document
    
    # Ensure path is a string
    path = str(path) if path else path
    topic_text = topic or Path(path).stem
    
    if doc_type == "assignment":
        doc = generate_assignment(topic_text, source_material=content)
        result = save_document(doc, Path(path).name, Path(path).parent)
        if result.get("status") == "success":
            return _build_result("Created Document", result.get("path", path), "success", result.get("message", ""), "Word Document", result)
        return _build_result("Create Document", path, "failed", result.get("message", ""), "Word Document")
    
    if doc_type == "research_proposal":
        doc = generate_research_proposal(topic_text, source_material=content)
        result = save_document(doc, Path(path).name, Path(path).parent)
        if result.get("status") == "success":
            return _build_result("Created Document", result.get("path", path), "success", result.get("message", ""), "Word Document", result)
        return _build_result("Create Document", path, "failed", result.get("message", ""), "Word Document")
    
    if doc_type == "report":
        doc = generate_report(topic_text, source_material=content)
        result = save_document(doc, Path(path).name, Path(path).parent)
        if result.get("status") == "success":
            return _build_result("Created Document", result.get("path", path), "success", result.get("message", ""), "Word Document", result)
        return _build_result("Create Document", path, "failed", result.get("message", ""), "Word Document")
    
    if doc_type == "schedule":
        # Extract person name from topic if present (e.g., "Leo's Training Schedule")
        import re
        person_name = None
        name_match = re.search(r"([A-Z][a-z]+)'s\s+(.+)", topic_text)
        if name_match:
            person_name = name_match.group(1)
            topic_text = name_match.group(2)
        doc = generate_schedule(topic_text, person_name)
        # Use the path that was passed in (already has correct filename)
        result = save_document(doc, Path(path).name, Path(path).parent)
        if result.get("status") == "success":
            return _build_result("Created Document", result.get("path", path), "success", result.get("message", ""), "Word Document", result)
        return _build_result("Create Document", path, "failed", result.get("message", ""), "Word Document")

    if doc_type == "dynamic_blueprint":
        doc = generate_dynamic_document(topic_text, instruction=instruction, source_material=content)
        result = save_document(doc, Path(path).name, Path(path).parent)
        if result.get("status") == "success":
            return _build_result("Created Document", result.get("path", path), "success", result.get("message", ""), "Word Document", result)
        return _build_result("Create Document", path, "failed", result.get("message", ""), "Word Document")

    if doc_type == "exam_paper":
        doc = generate_exam_document(topic_text, instruction=instruction, source_material=content)
        result = save_document(doc, Path(path).name, Path(path).parent)
        if result.get("status") == "success":
            return _build_result("Created Document", result.get("path", path), "success", result.get("message", ""), "Word Document", result)
        return _build_result("Create Document", path, "failed", result.get("message", ""), "Word Document")

    # Non-hardcoded DOCX fallback: generate a rich dynamic document
    if Path(path).suffix.lower() == ".docx":
        doc = generate_dynamic_document(topic_text, instruction=instruction, source_material=content)
        result = save_document(doc, Path(path).name, Path(path).parent)
        if result.get("status") == "success":
            return _build_result("Created Document", result.get("path", path), "success", result.get("message", ""), "Word Document", result)
        return _build_result("Create Document", path, "failed", result.get("message", ""), "Word Document")
    
    content = content or f"Draft document on {topic_text}.\n\nThis document was generated by the AMAZON assistant.\n"
    return create_document(path, content=content, topic=topic_text)


def _sanitize_filename(name):
    """Convert any string to a valid filename."""
    import re
    sanitized = re.sub(r'[\\/*?:"<>|]', "", name)
    sanitized = re.sub(r'\s+', '_', sanitized)
    return sanitized[:50].strip('_')
